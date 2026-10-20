"""Database-backed dashboard and live, versioned Grant Advisor reports."""
import json
from io import BytesIO
from uuid import uuid4
import httpx
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field, ConfigDict
try:
    from .database import Store
except ImportError:
    from database import Store

router = APIRouter(prefix='/api')
_store = None

def store():
    global _store
    if _store is None:
        _store = Store()
        _store.seed()
    return _store

class ApplicationEdit(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: int = Field(ge=1)
    ask: float | None = Field(default=None, gt=0, le=1000000, allow_inf_nan=False)
    cash_gap: float | None = Field(default=None, gt=0, le=1000000, allow_inf_nan=False)
    use: str | None = Field(default=None, min_length=1, max_length=3000)
    status: str | None = Field(default=None, pattern='^(ready|info|rev|new)$')
    notes: dict[str, str] | None = None
    fin: dict | None = None

class MemoRequest(BaseModel):
    applicant_ids: list[str] | None = None
    horizon_days: int = Field(default=90, ge=1, le=365)

@router.get('/bootstrap')
def bootstrap():
    return store().bootstrap()

@router.get('/applications/{id}')
def application(id: str):
    row = store().application(id)
    if row is None: raise HTTPException(404, 'Application not found')
    return row

@router.patch('/applications/{id}')
def edit(id: str, req: ApplicationEdit):
    changes=req.model_dump(exclude_none=True, exclude={'version'})
    if 'fin' in changes:
        existing=application(id)['fin']
        if set(changes['fin']) != set(existing): raise HTTPException(422,'Financial summary must preserve its original fields')
        for field in ('appRev','taxRev','quote'):
            value=changes['fin'][field]
            if not isinstance(value,(int,float)) or isinstance(value,bool) or not 0<=value<=1e12:
                raise HTTPException(422,'Invalid financial amount')
        quarters=changes['fin']['q']
        if not isinstance(quarters,list) or len(quarters)!=3 or any(v is not None and (not isinstance(v,(int,float)) or isinstance(v,bool) or not 0<=v<=1e12) for v in quarters):
            raise HTTPException(422,'Provide three quarterly amounts or null for missing values')
    try: return store().edit(id,changes,req.version)
    except KeyError: raise HTTPException(404,'Application not found')
    except ValueError as exc: raise HTTPException(409,str(exc))

async def invoke_agent(url, prompt):
    user='memo-'+str(uuid4())
    async with httpx.AsyncClient(timeout=httpx.Timeout(290,connect=10)) as client:
        response=await client.post(f'{url}/apps/grant_advisor/users/{user}/sessions',json={})
        response.raise_for_status()
        session=response.json()['id']
        response=await client.post(f'{url}/run',json={'appName':'grant_advisor','userId':user,'sessionId':session,
            'newMessage':{'role':'user','parts':[{'text':prompt}]}})
        response.raise_for_status()
        return response.json()

def extract_evidence(events,snapshot):
    reports=[];texts=[]
    for event in events:
        if event.get('errorCode'): raise ValueError('Agent execution failed')
        for p in (event.get('content') or {}).get('parts',[]):
            r=p.get('functionResponse') or {}
            if r.get('name')=='compare_applicants': reports.append(r['response'])
            if p.get('text','').strip() and not p.get('thought'): texts.append(p['text'].strip())
    if not reports or not texts: raise ValueError('Agent did not return comparison evidence and a memo')
    report=reports[-1]
    if report.get('status')!='ok': raise ValueError('Comparison did not complete')
    expected={a['business_id']:a for a in snapshot['applicants']}
    actual={a['business_id']:a for a in report['applicants']}
    if expected.keys()!=actual.keys() or report['budget']!=snapshot['budget'] or report['horizon_days']!=snapshot['horizon_days']:
        raise ValueError('Agent comparison does not match saved inputs')
    for id,a in expected.items():
        for key in ('requested_amount','cash_gap','input_type','evidence_note','operating_assumption'):
            if actual[id].get(key)!=a[key]: raise ValueError('Agent changed saved financial inputs')
    return report,'\n\n'.join(texts)

@router.post('/applications/{id}/memos')
async def generate_memo(id: str,req: MemoRequest):
    import os
    selected=application(id)
    data=store().bootstrap()
    ids=req.applicant_ids
    if ids is None:
        ids=[id]+[a['id'] for a in data['applications'] if a['id']!=id][:5]
    if not 2<=len(ids)<=6 or len(set(ids))!=len(ids) or id not in ids:
        raise HTTPException(422,'Choose 2–6 distinct applications including the selected application')
    by_id={a['id']:a for a in data['applications']}
    if any(x not in by_id for x in ids): raise HTTPException(404,'Application not found')
    records=[by_id[x] for x in ids]
    snapshot={'selected_application_id':id,'budget':data['settings']['round_budget'],'horizon_days':req.horizon_days,
        'application_versions':{a['id']:a['version'] for a in records},'applications':records,
        'applicants':[{'business_id':a['id'],'requested_amount':a['ask'],'cash_gap':a['cash_gap'],
            'input_type':a['input_type'],'evidence_note':a['evidence_note'],
            'operating_assumption':a['operating_assumption']} for a in records]}
    run_id=store().create_run(id,snapshot)
    queue=os.getenv('MEMO_TASK_QUEUE')
    if queue:
        try:
            from google.cloud import tasks_v2
            try: from .auth import task_token
            except ImportError: from auth import task_token
            client=tasks_v2.CloudTasksClient()
            task={'http_request':{'http_method':tasks_v2.HttpMethod.POST,
                'url':os.environ['PUBLIC_BACKEND_URL']+'/api/internal/memo-runs/'+run_id,
                'headers':{'Content-Type':'application/json','X-ReliefIQ-Task':task_token()},'body':b'{}'},
                'dispatch_deadline':{'seconds':600}}
            client.create_task(parent=queue,task=task)
            return {'run_id':run_id,'status':'running','mode':'queued'}
        except Exception as exc:
            store().fail(run_id,'Could not enqueue memo job: '+type(exc).__name__)
            raise HTTPException(503,'Could not queue memo generation') from exc
    if not store().claim_run(run_id): return {'run_id':run_id,'status':'running','mode':'queued'}
    return await execute_memo(run_id,id,snapshot)

async def execute_memo(run_id,id,snapshot):
    import os
    selected=next(a for a in snapshot['applications'] if a['id']==id)
    prompt=('Generate the complete seven-part dashboard GRANT MEMO for '+selected['name']+' ('+id+'). '
        'You MUST call compare_applicants with EXACTLY the applicants, budget and horizon below. '
        'The shortlist is a saved scenario, not all eight applications. Do not alter its inputs. '
        'Explain any gap exceeding the requested amount and any no-feasible-allocation result. '
        'Use exactly these seven numbered sections: 1. Which business, by name and address, and how much; '
        '2. What it protects—the cascade, named; 3. The traversal itself, shown; '
        '4. Why this one and not the runner-up; 5. What you are assuming; '
        '6. Which parts of the evidence are modeled; 7. Who the neighborhood is. '
        'When no funding is feasible, state the exact cash-gap feasibility reason; do not attribute it to zero graph exposure for every applicant. '
        'Financials, owners and document descriptions are simulated; never describe them as independently verified. '
        'Include sources, evidence-review limitations and conditional outcomes. Do not call other tools: compare_applicants already gathers specialist evidence. Saved input snapshot: '+json.dumps({k:snapshot[k] for k in ('selected_application_id','budget','horizon_days','applicants')}))
    try:
        events=await invoke_agent(os.getenv('ADK_AGENT_URL','https://cascade-analyst-89073885580.us-central1.run.app').rstrip('/'),prompt)
        evidence,content=extract_evidence(events,snapshot)
        memo=store().finish(run_id,id,evidence,content)
        return {**memo,'reply':content,'tool_calls':['compare_applicants'],'mode':'live','agent_app':'grant_advisor'}
    except (httpx.HTTPError,ValueError,KeyError,TypeError) as exc:
        store().fail(run_id,type(exc).__name__+': '+str(exc)[:500])
        raise HTTPException(502,{'message':'Live memo generation failed; no memo saved','run_id':run_id}) from exc

@router.get('/analysis-runs/{id}')
def run(id: str):
    row=store().run(id)
    if not row: raise HTTPException(404,'Analysis run not found')
    return row

@router.get('/applications/{id}/memos')
def memo_list(id: str):
    application(id)
    return [m for m in store().bootstrap()['memos'] if m['application_id']==id]

@router.get('/memos/{id}')
def memo(id: str):
    row=store().memo(id)
    if not row: raise HTTPException(404,'Memo not found')
    snapshot=run(row['run_id'])
    row['run']=snapshot
    row['stale']=any(application(aid)['version']!=version for aid,version in snapshot['input_snapshot']['application_versions'].items())
    return row

@router.get('/memos/{id}/download')
def download(id: str):
    from docx import Document
    row=memo(id)
    doc=Document();doc.add_heading('ReliefIQ — Grant Advisor Memo',0)
    doc.add_paragraph(f"Created {row['created_at']} | Application {row['application_id']} | Run {row['run_id']}")
    doc.add_paragraph('Simulated application inputs; live analytical evidence. Conditional decision support, subject to human review.')
    for line in row['content'].splitlines():
        if line.startswith('#'): doc.add_heading(line.lstrip('# ').replace('**',''),min(len(line)-len(line.lstrip('#')),3))
        else: doc.add_paragraph(line)
    doc.add_heading('Saved input versions',1)
    doc.add_paragraph(json.dumps(row['run']['input_snapshot']['application_versions'],indent=2))
    output=BytesIO();doc.save(output)
    return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition':f'attachment; filename="ReliefIQ-memo-{id}.docx"'})

class RecommendationRequest(BaseModel):
    recommended: bool

@router.post('/applications/{id}/recommendation')
def recommend(id: str, req: RecommendationRequest):
    row=application(id)
    try: return {'recommendations':store().recommend(id,row['ask'] if req.recommended else None),'status':'human_recommendation_not_disbursement'}
    except ValueError as exc: raise HTTPException(409,str(exc))

@router.post('/internal/memo-runs/{id}')
async def worker(id: str, request: Request):
    import hmac
    try: from .auth import task_token
    except ImportError: from auth import task_token
    if not hmac.compare_digest(request.headers.get('X-ReliefIQ-Task',''),task_token()):
        raise HTTPException(401,'Task authentication required')
    row=run(id)
    if not store().claim_run(id): return {'status':row['status']}
    return await execute_memo(id,row['application_id'],row['input_snapshot'])

@router.get('/health')
def health():
    return {'status':'ok','database':store().engine.dialect.name}

class BudgetEdit(BaseModel):
    round_budget: float = Field(gt=0,le=1000000,allow_inf_nan=False)

@router.patch('/funding-rounds/default')
def budget(req: BudgetEdit):
    try: return store().update_budget(req.round_budget)
    except ValueError as exc: raise HTTPException(409,str(exc))
